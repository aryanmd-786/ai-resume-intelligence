"""
services/resume_docx.py
-----------------------
Generate a professional DOCX resume using python-docx.
Template is driven by a theme dict.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def _safe(value: Any) -> str:
    return str(value or "").strip()


def _hex_to_rgb(hex_str: str) -> RGBColor:
    h = hex_str.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _set_run(run, *, size: int, bold: bool = False, italic: bool = False, color: RGBColor | None = None):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _bullet_lines(text: str) -> list[str]:
    lines: list[str] = []
    for raw in (text or "").splitlines():
        line = raw.strip().lstrip("•-*").strip()
        if line:
            lines.append(line)
    return lines


def _add_heading(doc: Document, text: str, theme: dict) -> None:
    label = text.upper() if theme.get("section_uppercase") else text
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(label)
    _set_run(
        run,
        size=int(theme.get("section_size", 11)) + 1,
        bold=True,
        color=_hex_to_rgb(theme["primary"]),
    )

    # Bottom border line
    pPr = p._p.get_or_add_pPr()
    borders = pPr.makeelement(qn("w:pBdr"), {})
    bottom = borders.makeelement(qn("w:bottom"), {})
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), theme["primary"].lstrip("#"))
    borders.append(bottom)
    pPr.append(borders)


def _add_body(doc: Document, text: str, theme: dict) -> None:
    p = doc.add_paragraph()
    run = p.add_run(_safe(text))
    _set_run(run, size=int(theme.get("body_size", 10)), color=_hex_to_rgb(theme["text"]))


def _add_bullet(doc: Document, text: str, theme: dict) -> None:
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(_safe(text))
    _set_run(run, size=int(theme.get("body_size", 10)), color=_hex_to_rgb(theme["text"]))


def generate_docx(data: dict, theme: dict, output_path: Path) -> Path:
    """Write a DOCX resume to output_path and return it."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()

    # Set default font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(int(theme.get("body_size", 10)))
    style.font.color.rgb = _hex_to_rgb(theme["text"])

    header_align = (
        WD_ALIGN_PARAGRAPH.CENTER
        if theme.get("header_align") == "center"
        else WD_ALIGN_PARAGRAPH.LEFT
    )

    # ------------------------------------------------------------------
    # Header
    # ------------------------------------------------------------------
    personal = data.get("personal", {}) or {}
    name = _safe(personal.get("full_name")) or "Your Name"

    p = doc.add_paragraph()
    p.alignment = header_align
    run = p.add_run(name)
    _set_run(
        run,
        size=int(theme.get("name_size", 20)),
        bold=True,
        color=_hex_to_rgb(theme["text"]),
    )

    contact_parts = [v for v in (personal.get("email"), personal.get("phone"), personal.get("location")) if _safe(v)]
    if contact_parts:
        p = doc.add_paragraph()
        p.alignment = header_align
        run = p.add_run("  •  ".join(_safe(x) for x in contact_parts))
        _set_run(run, size=int(theme.get("body_size", 10)) - 1, color=_hex_to_rgb(theme["muted"]))

    links = [v for v in (personal.get("linkedin"), personal.get("github"), personal.get("portfolio")) if _safe(v)]
    if links:
        p = doc.add_paragraph()
        p.alignment = header_align
        run = p.add_run("  |  ".join(_safe(x) for x in links))
        _set_run(run, size=int(theme.get("body_size", 10)) - 1, color=_hex_to_rgb(theme["muted"]))

    # ------------------------------------------------------------------
    # Summary
    # ------------------------------------------------------------------
    summary = _safe(data.get("summary"))
    if summary:
        _add_heading(doc, "Professional Summary", theme)
        _add_body(doc, summary, theme)

    # ------------------------------------------------------------------
    # Skills
    # ------------------------------------------------------------------
    skills = data.get("skills") or []
    if isinstance(skills, str):
        skills = [s.strip() for s in skills.split(",") if s.strip()]
    skills = [str(s).strip() for s in skills if str(s).strip()]
    if skills:
        _add_heading(doc, "Skills", theme)
        _add_body(doc, ", ".join(skills), theme)

    # ------------------------------------------------------------------
    # Experience
    # ------------------------------------------------------------------
    experience = data.get("experience") or []
    valid_exp = [e for e in experience if any(_safe(e.get(k)) for k in ("role", "company", "responsibilities"))]
    if valid_exp:
        _add_heading(doc, "Experience", theme)
        for e in valid_exp:
            role = _safe(e.get("role"))
            company = _safe(e.get("company"))
            duration = _safe(e.get("duration"))

            title = " — ".join(x for x in (role, company) if x) or "Experience"
            p = doc.add_paragraph()
            run = p.add_run(title)
            _set_run(run, size=int(theme.get("body_size", 10)) + 1, bold=True, color=_hex_to_rgb(theme["text"]))

            if duration:
                p = doc.add_paragraph()
                run = p.add_run(duration)
                _set_run(run, size=int(theme.get("body_size", 10)) - 1, italic=True, color=_hex_to_rgb(theme["muted"]))

            for b in _bullet_lines(e.get("responsibilities", "")):
                _add_bullet(doc, b, theme)

    # ------------------------------------------------------------------
    # Projects
    # ------------------------------------------------------------------
    projects = data.get("projects") or []
    valid_prj = [p for p in projects if any(_safe(p.get(k)) for k in ("name", "description"))]
    if valid_prj:
        _add_heading(doc, "Projects", theme)
        for p_ in valid_prj:
            name_p = _safe(p_.get("name")) or "Project"
            tech = _safe(p_.get("technologies"))
            url = _safe(p_.get("url"))

            para = doc.add_paragraph()
            run = para.add_run(name_p)
            _set_run(run, size=int(theme.get("body_size", 10)) + 1, bold=True, color=_hex_to_rgb(theme["text"]))

            meta_parts = []
            if tech:
                meta_parts.append(f"Tech: {tech}")
            if url:
                meta_parts.append(url)
            if meta_parts:
                para = doc.add_paragraph()
                run = para.add_run("  •  ".join(meta_parts))
                _set_run(run, size=int(theme.get("body_size", 10)) - 1, italic=True, color=_hex_to_rgb(theme["muted"]))

            for b in _bullet_lines(p_.get("description", "")):
                _add_bullet(doc, b, theme)

    # ------------------------------------------------------------------
    # Education
    # ------------------------------------------------------------------
    education = data.get("education") or []
    valid_edu = [e for e in education if any(_safe(e.get(k)) for k in ("degree", "institution"))]
    if valid_edu:
        _add_heading(doc, "Education", theme)
        for e in valid_edu:
            degree = _safe(e.get("degree"))
            inst = _safe(e.get("institution"))
            year = _safe(e.get("year"))
            score = _safe(e.get("score"))

            title = " — ".join(x for x in (degree, inst) if x) or "Education"
            para = doc.add_paragraph()
            run = para.add_run(title)
            _set_run(run, size=int(theme.get("body_size", 10)) + 1, bold=True, color=_hex_to_rgb(theme["text"]))

            meta_parts = []
            if year:
                meta_parts.append(year)
            if score:
                meta_parts.append(f"Score: {score}")
            if meta_parts:
                para = doc.add_paragraph()
                run = para.add_run("  •  ".join(meta_parts))
                _set_run(run, size=int(theme.get("body_size", 10)) - 1, italic=True, color=_hex_to_rgb(theme["muted"]))

    # ------------------------------------------------------------------
    # Certifications / Achievements
    # ------------------------------------------------------------------
    certs = [c for c in (data.get("certifications") or []) if _safe(c)]
    if certs:
        _add_heading(doc, "Certifications", theme)
        for c in certs:
            _add_bullet(doc, _safe(c), theme)

    achievements = [a for a in (data.get("achievements") or []) if _safe(a)]
    if achievements:
        _add_heading(doc, "Achievements", theme)
        for a in achievements:
            _add_bullet(doc, _safe(a), theme)

    # ------------------------------------------------------------------
    # Languages / Interests
    # ------------------------------------------------------------------
    languages = data.get("languages") or []
    if isinstance(languages, str):
        languages = [s.strip() for s in languages.split(",") if s.strip()]
    languages = [str(x).strip() for x in languages if str(x).strip()]
    if languages:
        _add_heading(doc, "Languages", theme)
        _add_body(doc, ", ".join(languages), theme)

    interests = data.get("interests") or []
    if isinstance(interests, str):
        interests = [s.strip() for s in interests.split(",") if s.strip()]
    interests = [str(x).strip() for x in interests if str(x).strip()]
    if interests:
        _add_heading(doc, "Interests", theme)
        _add_body(doc, ", ".join(interests), theme)

    doc.save(str(output_path))
    return output_path