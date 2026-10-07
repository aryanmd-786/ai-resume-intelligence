"""
services/resume_parser.py
-------------------------
Extract plain text from PDF and DOCX resumes.

Design goals:
- No crashes on malformed files.
- Clear, user-friendly error messages.
- Reasonable text cleanup so AI gets clean input.
"""

from __future__ import annotations

import io
import re
from pathlib import Path


class ResumeParserError(Exception):
    """Raised for any parsing-related failure."""


# ---------------------------------------------------------------------------
# Text cleanup
# ---------------------------------------------------------------------------
_MULTI_SPACE_RE = re.compile(r"[ \t]+")
_MULTI_NEWLINE_RE = re.compile(r"\n{3,}")
_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _clean_text(text: str) -> str:
    """Normalize whitespace and strip control characters."""
    if not text:
        return ""
    text = _CONTROL_RE.sub("", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _MULTI_SPACE_RE.sub(" ", text)
    text = _MULTI_NEWLINE_RE.sub("\n\n", text)
    return text.strip()


# ---------------------------------------------------------------------------
# PDF extraction (PyMuPDF / fitz)
# ---------------------------------------------------------------------------
def extract_text_from_pdf(data: bytes) -> str:
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:
        raise ResumeParserError(
            "PyMuPDF is not installed. Run: pip install PyMuPDF"
        ) from exc

    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise ResumeParserError(
            "Could not open this PDF. It may be corrupted or password-protected."
        ) from exc

    try:
        if doc.page_count == 0:
            raise ResumeParserError("PDF contains no pages.")

        pages_text: list[str] = []
        for page in doc:
            pages_text.append(page.get_text("text") or "")

        combined = "\n".join(pages_text)
    finally:
        doc.close()

    cleaned = _clean_text(combined)
    if len(cleaned) < 50:
        raise ResumeParserError(
            "No readable text found in this PDF. "
            "It may be a scanned image — try exporting a text-based PDF."
        )
    return cleaned


# ---------------------------------------------------------------------------
# DOCX extraction (python-docx)
# ---------------------------------------------------------------------------
def extract_text_from_docx(data: bytes) -> str:
    try:
        from docx import Document  # python-docx
    except ImportError as exc:
        raise ResumeParserError(
            "python-docx is not installed. Run: pip install python-docx"
        ) from exc

    try:
        stream = io.BytesIO(data)
        document = Document(stream)
    except Exception as exc:
        raise ResumeParserError(
            "Could not open this DOCX. It may be corrupted."
        ) from exc

    chunks: list[str] = []

    # Paragraphs
    for para in document.paragraphs:
        text = (para.text or "").strip()
        if text:
            chunks.append(text)

    # Tables (many resumes use tables for layout)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))

    combined = "\n".join(chunks)
    cleaned = _clean_text(combined)

    if len(cleaned) < 50:
        raise ResumeParserError(
            "No readable text found in this DOCX."
        )
    return cleaned


# ---------------------------------------------------------------------------
# Router
# ---------------------------------------------------------------------------
def parse_resume(data: bytes, filename: str) -> str:
    """
    Route to the correct extractor based on file extension.

    Raises ResumeParserError on any failure.
    """
    if not data:
        raise ResumeParserError("File is empty.")

    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(data)
    if ext == ".docx":
        return extract_text_from_docx(data)

    raise ResumeParserError(
        f"Unsupported file type '{ext}'. Only .pdf and .docx are supported."
    )